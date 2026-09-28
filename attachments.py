"""ATTACHMENTS. Turn an attached file into text for the prompt: csv/txt/md/json as they are, pdf via pypdf,
xlsx via openpyxl (every sheet, one line per row, cells joined by tabs), docx via python-docx (paragraphs,
then tables). Anything else -> die naming the file and its size: never guess a file's contents.

A cap (MAX_CHARS, or --max-attachment-chars on run.py) stops the run with a clear message instead of
silently feeding a huge file to every model.
"""
from __future__ import annotations

from pathlib import Path
from typing import NoReturn

TEXT = {".csv", ".txt", ".md", ".json"}
CONVERTED = {".pdf", ".xlsx", ".docx"}
MAX_CHARS = 200_000  # about 50k tokens; one file over this needs a decision, not a default


def die(msg: str) -> NoReturn:
    raise SystemExit(f"ERROR: {msg}")


def read_attachment(path: str | Path, max_chars: int = MAX_CHARS) -> str:
    """File -> text. Unknown type or over the cap -> die with the file name and size."""
    path = Path(path)
    if not path.is_file():
        die(f"attachment not found: {path}")
    ext = path.suffix.lower()
    if ext in TEXT:
        text = path.read_text(encoding="utf-8-sig").rstrip()  # utf-8-sig drops the BOM some APEX CSVs carry
    elif ext == ".pdf":
        text = _pdf(path)
    elif ext == ".xlsx":
        text = _xlsx(path)
    elif ext == ".docx":
        text = _docx(path)
    else:
        die(f"attachment {path.name} is {ext or 'no extension'} ({path.stat().st_size:,} bytes); "
            f"supported: {sorted(TEXT | CONVERTED)}. Stop and decide, do not guess its contents.")
    if len(text) > max_chars:
        die(f"attachment {path.name} is {len(text):,} chars, cap is {max_chars:,}; "
            f"trim it or raise --max-attachment-chars on purpose")
    return text


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    pages = [(p.extract_text() or "").rstrip() for p in PdfReader(str(path)).pages]
    text = "\n\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1))
    if not any(pages):
        die(f"attachment {path.name}: {len(pages)} page(s), no extractable text (scanned?). Stop and decide.")
    return text


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)  # data_only: cell values, not formulas
    out = []
    for ws in wb.worksheets:
        lines = ["\t".join("" if c is None else str(c) for c in row) for row in ws.iter_rows(values_only=True)]
        out.append(f"--- sheet {ws.title} ---\n" + "\n".join(lines).rstrip())
    return "\n\n".join(out)


def _docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for i, t in enumerate(d.tables, 1):
        rows = ["\t".join(c.text.strip() for c in r.cells) for r in t.rows]
        parts.append(f"--- table {i} ---\n" + "\n".join(rows))
    return "\n".join(parts).rstrip()


def report(tasks: list[dict]) -> str:
    """Per-task attachment character counts, largest first, for the run dashboard."""
    rows = [(sum(t["metadata"].get("attachment_chars", {}).values()), t["id"], t["metadata"].get("attachment_chars", {}))
            for t in tasks if t["metadata"].get("attachment_chars")]
    if not rows:
        return "attachments  : none"
    rows.sort(reverse=True)
    total = sum(r[0] for r in rows)
    head = f"attachments  : {len(rows)} task(s) with files, {total:,} chars total, largest task {rows[0][1]} ({rows[0][0]:,})"
    lines = [f"  task {tid}: {n:,} chars  " + ", ".join(f"{k} {v:,}" for k, v in files.items()) for n, tid, files in rows[:10]]
    return "\n".join([head, *lines])
