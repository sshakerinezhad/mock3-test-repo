import pytest

import attachments
import run


def test_text_types_read_as_is(tmp_path):
    (tmp_path / "a.csv").write_bytes(b"\xef\xbb\xbfx,y\r\n1,2\r\n")
    assert attachments.read_attachment(tmp_path / "a.csv") == "x,y\n1,2"


def test_xlsx_every_sheet_rows_as_tabs(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["item", "qty"])
    ws.append(["pen", 3])
    wb.create_sheet("Notes").append(["hello"])
    wb.save(tmp_path / "s.xlsx")
    assert attachments.read_attachment(tmp_path / "s.xlsx") == "--- sheet Sales ---\nitem\tqty\npen\t3\n\n--- sheet Notes ---\nhello"


def test_docx_paragraphs_then_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Contract terms")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "fee", "100"
    d.save(tmp_path / "c.docx")
    assert attachments.read_attachment(tmp_path / "c.docx") == "Contract terms\n--- table 1 ---\nfee\t100"


def test_pdf_without_text_dies_instead_of_guessing(tmp_path):
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with open(tmp_path / "scan.pdf", "wb") as f:
        w.write(f)
    with pytest.raises(SystemExit, match="no extractable text"):
        attachments.read_attachment(tmp_path / "scan.pdf")


def test_unknown_type_dies_with_name_and_size(tmp_path):
    (tmp_path / "pic.png").write_bytes(b"\x89PNG" + b"0" * 96)
    with pytest.raises(SystemExit, match=r"pic.png is .png \(100 bytes\)"):
        attachments.read_attachment(tmp_path / "pic.png")


def test_cap_dies_clearly(tmp_path):
    (tmp_path / "big.txt").write_text("x" * 50)
    with pytest.raises(SystemExit, match="50 chars, cap is 10"):
        attachments.read_attachment(tmp_path / "big.txt", max_chars=10)
    assert len(attachments.read_attachment(tmp_path / "big.txt", max_chars=50)) == 50


def test_report_lists_largest_first():
    tasks = [{"id": 1, "metadata": {"attachment_chars": {"a.csv": 10}}},
             {"id": 2, "metadata": {"attachment_chars": {"b.pdf": 500, "c.csv": 5}}},
             {"id": 3, "metadata": {}}]
    r = attachments.report(tasks)
    assert r.startswith("attachments  : 2 task(s) with files, 515 chars total, largest task 2 (505)")
    assert r.splitlines()[1].startswith("  task 2: 505 chars")
    assert attachments.report([{"id": 9, "metadata": {}}]) == "attachments  : none"


def test_apex_loader_names_the_task_on_a_bad_attachment(tmp_path):
    d = tmp_path / "documents" / "3"
    d.mkdir(parents=True)
    (d / "img.png").write_bytes(b"0" * 10)
    p = tmp_path / "train.csv"
    p.write_text('Task ID,Domain,Prompt,File Attachments\n3,Legal,"look","documents/3/img.png"\n')
    with pytest.raises(SystemExit, match=r"task 3: attachment img.png is .png"):
        run.load_tasks(p)
